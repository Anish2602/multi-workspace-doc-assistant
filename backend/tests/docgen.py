"""Build small PDF/DOCX files in-memory for tests."""

import io

from docx import Document
from fpdf import FPDF


def make_pdf(pages: list[str]) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for text in pages:
        pdf.add_page()
        pdf.multi_cell(0, 6, text)
    return bytes(pdf.output())


def make_docx(sections: dict[str, list[str]]) -> bytes:
    doc = Document()
    for heading, paragraphs in sections.items():
        doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
