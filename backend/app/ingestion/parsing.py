"""Turn uploaded bytes into text blocks that remember where they came from.

Each `Block` is roughly a paragraph, tagged with its page (PDF) and the nearest
heading (DOCX/Markdown) so citations can point at "page 3" or "§ Leave policy".
"""

import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


class UnsupportedDocument(ValueError):
    """The upload can't be ingested (wrong type, corrupt, or no extractable text)."""


@dataclass(frozen=True)
class Block:
    text: str
    page: int | None = None
    section: str | None = None


def parse_document(filename: str, data: bytes) -> list[Block]:
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocument(
            f"Unsupported file type '{ext or filename}'. Use PDF, DOCX, MD or TXT."
        )

    if ext == ".pdf":
        blocks = _parse_pdf(data)
    elif ext == ".docx":
        blocks = _parse_docx(data)
    else:
        blocks = _parse_text(_decode(data), markdown=ext in {".md", ".markdown"})

    blocks = [b for b in blocks if b.text.strip()]
    if not blocks:
        raise UnsupportedDocument(
            "No extractable text found (scanned PDFs/images aren't supported)."
        )
    return blocks


def _decode(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise UnsupportedDocument("File looks binary, not text.")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def _parse_pdf(data: bytes) -> list[Block]:
    if not data.startswith(b"%PDF"):
        raise UnsupportedDocument("File has a .pdf extension but isn't a PDF.")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnsupportedDocument("Encrypted PDFs aren't supported.")
        blocks = []
        for page_no, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            blocks += [Block(p, page=page_no) for p in _paragraphs(text)]
        return blocks
    except PdfReadError as exc:
        raise UnsupportedDocument("Could not read this PDF (corrupt?).") from exc


def _parse_docx(data: bytes) -> list[Block]:
    if not data.startswith(b"PK"):
        raise UnsupportedDocument("File has a .docx extension but isn't a DOCX.")
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of zip/xml errors
        raise UnsupportedDocument("Could not read this DOCX (corrupt?).") from exc

    blocks: list[Block] = []
    section: str | None = None
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            section = text[:255]
            continue
        blocks.append(Block(text, section=section))
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(Block(" | ".join(cells), section=section))
    return blocks


_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")


def _parse_text(text: str, *, markdown: bool) -> list[Block]:
    if not markdown:
        return [Block(p) for p in _paragraphs(text)]

    blocks: list[Block] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.extend(Block(p, section=section) for p in _paragraphs("\n".join(buffer)))
            buffer.clear()

    for line in text.splitlines():
        match = _MD_HEADING.match(line)
        if match:
            flush()
            section = match.group(1)[:255] or section
        else:
            buffer.append(line)
    flush()
    return blocks
